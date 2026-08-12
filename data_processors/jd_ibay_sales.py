# -*- coding: utf-8 -*-
"""
京东艾贝自营仓销售出库 - 数据清洗模块

功能：清洗原始销售数据 → 分类蓝单/红单 → 匹配物料 → 生成上传模板
使用：由中台调度系统统一调用
"""
import os
import re
import calendar
from datetime import datetime, timedelta
from difflib import SequenceMatcher
import pandas as pd
from loguru import logger


# ==================== 配置（可被用户覆盖） ====================

DEFAULT_CONFIG = {
    "material_path": r"D:\ly\京东艾贝自营出库\新物料.xlsx",
    "output_dir": r"D:\ly\京东艾贝自营出库",
    "similarity_threshold": 0.92,
}

FIXED_VALUES = {
    '购货单位': '京东-艾贝自营店',
    '销售方式': '赊销',
    '发货仓库': '京东自营仓'
}


# ==================== 主入口（5步流程） ====================

def process(input_file: str, context: dict = None, params: dict = None, **kwargs) -> str:
    """
    清洗京东艾贝自营仓销售出库数据

    Args:
        input_file: 原始Excel文件路径
        context: 任务上下文（包含 user_params）
        params: 运行时参数

    Returns:
        清洗后的文件路径
    """
    # 1. 加载配置
    config = _get_config(context, params)
    logger.info(f"[清洗] 开始处理: {input_file}")

    # 2. 读取并清洗数据
    df = _read_and_clean(input_file)
    if df is None or len(df) == 0:
        logger.warning("[清洗] 无有效数据")
        return _save_empty_result(config['output_dir'])

    # 3. 分类汇总
    blue, red = _classify(df)
    logger.info(f"[清洗] 蓝单 {len(blue)} 个SKU，红单 {len(red)} 个SKU")

    # 4. 加载物料表
    name_map, unit_map = _load_material(config['material_path'])

    # 5. 匹配物料并生成上传文件
    output_file = _generate_output(blue, red, name_map, unit_map, config)

    logger.info(f"[清洗] 完成: {output_file}")
    return output_file


# ==================== 配置加载 ====================

def _get_config(context, params):
    """合并配置：params > user_params > 默认"""
    user_params = context.get('user_params', {}) if context else {}
    config = DEFAULT_CONFIG.copy()
    config.update(user_params)
    if params:
        config.update({k: v for k, v in params.items() if v is not None})
    return config


# ==================== 数据读取和清洗 ====================

def _read_and_clean(input_file):
    """读取Excel并清洗"""
    df = pd.read_excel(input_file)
    logger.info(f"[清洗] 原始数据: {len(df)} 行")

    # 列映射（京东导出格式固定：前12列）
    if len(df.columns) >= 12:
        df.columns = ['单据类型', '单据编号', '采购单号', '订单号',
                      '采销员', '合同主体', 'SKU编码', '数量',
                      '金额', '币种', '业务日期', '采购渠道'] + list(df.columns[12:])

    # 去空行
    df = df.dropna(how='all')

    # SKU清洗
    if 'SKU编码' in df.columns:
        df['SKU编码'] = df['SKU编码'].astype(str).str.strip()
        df['SKU编码'] = df['SKU编码'].apply(lambda x: x.split('.')[0] if '.' in x else x)
        df = df[~df['SKU编码'].isin(['', 'nan', 'None'])]

    # 数量清洗
    if '数量' in df.columns:
        df['数量'] = pd.to_numeric(df['数量'], errors='coerce')
        df = df[df['数量'].notna() & (df['数量'] != 0)]

    # 日期清洗
    if '业务日期' in df.columns:
        df['业务日期'] = pd.to_datetime(df['业务日期'], errors='coerce')

    return df


# ==================== 分类汇总 ====================

def _classify(df):
    """
    分类蓝单/红单，按SKU汇总
    蓝单 = 销售出库（数量为负）
    红单 = 退货入库（数量为正）
    """
    df['数量绝对值'] = df['数量'].abs()

    # 蓝单
    blue = df[df['数量'] < 0].groupby('SKU编码')['数量绝对值'].sum().reset_index()
    blue.columns = ['SKU编码', '蓝单数量']

    # 红单
    red = df[df['数量'] > 0].groupby('SKU编码')['数量绝对值'].sum().reset_index()
    red.columns = ['SKU编码', '红单数量']

    return blue, red


# ==================== 物料匹配 ====================

def _load_material(file_path):
    """加载物料表"""
    if not os.path.exists(file_path):
        logger.warning(f"[清洗] 物料表不存在: {file_path}")
        return {}, {}

    df = pd.read_excel(file_path, engine='openpyxl')
    name_map, unit_map = {}, {}

    for _, row in df.iterrows():
        sku = str(row.iloc[0]).strip()
        name = str(row.iloc[1]).strip() if pd.notna(row.iloc[1]) else ''
        unit = str(row.iloc[2]).strip() if len(row) > 2 and pd.notna(row.iloc[2]) else '个'
        if sku and name:
            name_map[sku] = name
            unit_map[sku] = unit

    logger.info(f"[清洗] 物料表加载: {len(name_map)} 条")
    return name_map, unit_map


def _match_sku(sku, name_map, unit_map, threshold, unmatched_log):
    """
    匹配SKU（精确 → 去前导零 → 补前导零 → 模糊匹配）
    返回: (匹配到的SKU, 物料名称, 单位)
    """
    # 精确匹配
    if sku in name_map:
        return sku, name_map[sku], unit_map.get(sku, '个')

    # 去掉前导零
    stripped = sku.lstrip('0')
    if stripped in name_map:
        return stripped, name_map[stripped], unit_map.get(stripped, '个')

    # 补前导零到13位
    padded = sku.zfill(13)
    if padded in name_map:
        return padded, name_map[padded], unit_map.get(padded, '个')

    # 模糊匹配
    best = None
    best_ratio = 0
    for key in name_map.keys():
        ratio = SequenceMatcher(None, sku, key).ratio()
        if ratio > best_ratio:
            best_ratio = ratio
            best = key

    if best and best_ratio >= threshold:
        logger.debug(f"[模糊匹配] {sku} → {best} ({best_ratio:.0%})")
        return best, name_map[best], unit_map.get(best, '个')

    # 未匹配
    unmatched_log.append(sku)
    return sku, sku, ''


# ==================== 日期处理 ====================

def _parse_date(value):
    """把各种格式转成日期"""
    if value is None or isinstance(value, datetime):
        return value

    s = str(value).strip()
    if not s or s in ('nan', 'None', 'NaT', ''):
        return None

    s = s.split(' ')[0].split('T')[0]

    m = re.match(r'(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})', s)
    if m:
        try:
            return datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            pass

    if s.isdigit() and len(s) == 8:
        try:
            return datetime(int(s[:4]), int(s[4:6]), int(s[6:8]))
        except ValueError:
            pass

    return None


def _get_doc_date(df):
    """从数据中取出现最多的月份，返回该月最后一天"""
    dates = df['业务日期'].dropna() if '业务日期' in df.columns else []

    if len(dates) > 0:
        months = dates.apply(lambda x: (x.year, x.month))
        best = months.value_counts().idxmax()
        year, month = best
        last_day = calendar.monthrange(year, month)[1]
        return datetime(year, month, last_day).strftime('%Y-%m-%d')

    # 兜底
    today = datetime.now()
    if today.day <= 10:
        return (today.replace(day=1) - timedelta(days=1)).strftime('%Y-%m-%d')
    return today.strftime('%Y-%m-%d')


# ==================== 生成输出 ====================

def _generate_output(blue, red, name_map, unit_map, config):
    """生成上传模板文件"""
    threshold = config['similarity_threshold']
    output_dir = config['output_dir']
    os.makedirs(output_dir, exist_ok=True)

    unmatched_log = []

    # 表头（27列）
    headers = [
        '[表头]单据编号(*)', '[表头]单据日期(*)', '[表头]购货单位(*)', '[表头]销售方式(*)',
        '[表头]摘要', '[表头]交货地点', '[表头]联系人', '[表头]联系人电话',
        '[表头]收货地址', '[表头]红蓝单', '物料代码(*)', '物料名称',
        '规格型号', '单位(*)', '发货仓库(*)', '实发数量(*)', '单位成本(*)', '备注',
        '辅助属性', '客户料号', '客户商品名称', '生产/采购日期', '保质期(天)',
        '有效期至', '批号', '销售单价', '折扣率'
    ]

    # 生成单据编号
    blue_no = datetime.now().strftime('%Y%m%d%H%M%S')
    red_no = (datetime.now() + timedelta(seconds=100)).strftime('%Y%m%d%H%M%S')
    doc_date = _get_doc_date(pd.concat([blue, red], axis=0) if len(blue) > 0 or len(red) > 0 else pd.DataFrame())

    # 构建数据行
    rows = []

    # 蓝单
    first = True
    for _, row in blue.iterrows():
        sku, name, unit = _match_sku(row['SKU编码'], name_map, unit_map, threshold, unmatched_log)
        rows.append({
            '单据编号': blue_no if first else '',
            '单据日期': doc_date if first else '',
            '购货单位': FIXED_VALUES['购货单位'] if first else '',
            '销售方式': FIXED_VALUES['销售方式'] if first else '',
            '红蓝单': '蓝单' if first else '',
            '物料代码': sku,
            '物料名称': name,
            '单位': unit,
            '实发数量': row['蓝单数量'],
        })
        first = False

    # 红单
    first = True
    for _, row in red.iterrows():
        sku, name, unit = _match_sku(row['SKU编码'], name_map, unit_map, threshold, unmatched_log)
        rows.append({
            '单据编号': red_no if first else '',
            '单据日期': doc_date if first else '',
            '购货单位': FIXED_VALUES['购货单位'] if first else '',
            '销售方式': FIXED_VALUES['销售方式'] if first else '',
            '红蓝单': '红单' if first else '',
            '物料代码': sku,
            '物料名称': name,
            '单位': unit,
            '实发数量': row['红单数量'],
        })
        first = False

    # 写入Excel
    data = [headers]
    for r in rows:
        row = [''] * len(headers)
        row[0] = r['单据编号']
        row[1] = r['单据日期']
        row[2] = r['购货单位']
        row[3] = r['销售方式']
        row[9] = r['红蓝单']
        row[10] = r['物料代码']
        row[11] = r['物料名称']
        row[13] = r['单位']
        row[14] = FIXED_VALUES['发货仓库']
        row[15] = r['实发数量']
        row[22] = 0
        row[25] = 1
        row[26] = 0
        data.append(row)

    # 保存文件
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = os.path.join(output_dir, f"京东艾贝自营出库单汇总_{timestamp}.xlsx")

    pd.DataFrame(data).to_excel(output_file, index=False, header=False, sheet_name='销售出库单', engine='openpyxl')

    # 加边框
    _apply_border(output_file, len(rows))

    # 记录未匹配日志
    if unmatched_log:
        log_file = os.path.join(output_dir, f"未匹配物料_{datetime.now().strftime('%Y%m%d')}.txt")
        with open(log_file, 'w', encoding='utf-8') as f:
            f.write(f"未匹配物料 ({len(unmatched_log)} 个)\n")
            f.write("=" * 40 + "\n")
            for sku in sorted(set(unmatched_log)):
                f.write(f"  {sku}\n")
        logger.info(f"[清洗] 未匹配物料: {log_file}")

    return output_file


def _apply_border(filepath, row_count):
    """加边框"""
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Border, Side, Alignment

        wb = load_workbook(filepath)
        ws = wb.active
        thin = Border(left=Side(style='thin'), right=Side(style='thin'),
                      top=Side(style='thin'), bottom=Side(style='thin'))

        for row in ws.iter_rows(max_row=row_count + 1, max_col=27):
            for cell in row:
                cell.border = thin
                cell.alignment = Alignment(horizontal='left' if cell.row > 1 else 'center', vertical='center')

        wb.save(filepath)
    except Exception:
        pass  # 边框只是美化，失败不影响


def _save_empty_result(output_dir):
    """无数据时生成空结果"""
    os.makedirs(output_dir, exist_ok=True)
    output_file = os.path.join(output_dir, f"无数据_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    pd.DataFrame([{'状态': '无有效数据', '时间': datetime.now()}]).to_excel(output_file, index=False)
    return output_file