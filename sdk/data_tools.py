# -*- coding: utf-8 -*-
"""
通用数据清洗工具库

把 data_clean 中反复出现的通用操作（读 Excel、SKU 匹配、分类汇总、Excel 美化）
抽成可复用函数，让每个 data_clean 业务文件只写差异化的匹配规则。

借鉴 dj_new_server 的 Utils 思想：把"每个需求都要写一遍的样板代码"收敛到工具层。
"""
import os
import re
import calendar
from datetime import datetime, timedelta
from difflib import SequenceMatcher
from typing import Any, Callable, Dict, List, Optional, Tuple
from loguru import logger

import pandas as pd


# ==================== 智能读取 ====================

def read_excel_smart(
    file_path: str,
    header_keywords: List[str] = None,
    header_search_rows: int = 20,
) -> pd.DataFrame:
    """
    智能读取 Excel/CSV，自动检测表头行。

    Args:
        file_path: 文件路径（支持 .xlsx / .xls / .csv）
        header_keywords: 表头行关键词列表，命中任一即认为是表头行
        header_search_rows: 在前 N 行内搜索表头

    Returns:
        清洗后的 DataFrame（表头已设置，空行空列已删除）
    """
    ext = os.path.splitext(file_path)[1].lower()
    keywords = header_keywords or ["日期", "时间", "金额", "数量", "SKU", "订单号", "名称"]

    def _find_header(df_raw: pd.DataFrame) -> int:
        for i in range(min(header_search_rows, len(df_raw))):
            row_vals = [str(v).strip() for v in df_raw.iloc[i].tolist() if str(v).strip() != "nan"]
            joined = "|".join(row_vals)
            if any(kw in joined for kw in keywords):
                return i
        return 0

    if ext == ".csv":
        for enc in ["utf-8-sig", "utf-8", "gbk", "gb18030"]:
            try:
                raw = pd.read_csv(file_path, header=None, dtype=str, encoding=enc)
                header_idx = _find_header(raw)
                df = pd.read_csv(file_path, header=header_idx, dtype=str, encoding=enc)
                logger.info(f"[data_tools] CSV读取(encoding={enc})，表头行: 第{header_idx}行")
                return df.dropna(axis=0, how="all").dropna(axis=1, how="all")
            except Exception:
                continue
        raise RuntimeError(f"CSV文件读取失败: {file_path}")
    else:
        raw = pd.read_excel(file_path, header=None, dtype=str)
        header_idx = _find_header(raw)
        logger.info(f"[data_tools] 表头行: 第{header_idx}行")
        df = pd.read_excel(file_path, header=header_idx, dtype=str)
        return df.dropna(axis=0, how="all").dropna(axis=1, how="all")


# ==================== 列名工具 ====================

def normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    """列名规范化：去空格、去换行"""
    df.columns = [str(c).replace("\n", "").replace(" ", "").strip() for c in df.columns]
    return df


def find_column(df: pd.DataFrame, keywords: List[str]) -> Optional[str]:
    """在列名中搜索关键词，返回第一个匹配的列名"""
    for kw in keywords:
        for col in df.columns:
            if kw in str(col):
                return col
    return None


# ==================== 数据清洗 ====================

def clean_sku_column(df: pd.DataFrame, col: str = "SKU编码") -> pd.DataFrame:
    """清洗 SKU 列：去空格、去小数点后缀、去空值"""
    if col not in df.columns:
        return df
    df[col] = df[col].astype(str).str.strip()
    df[col] = df[col].apply(lambda x: x.split(".")[0] if "." in x else x)
    df = df[~df[col].isin(["", "nan", "None"])]
    return df


def clean_numeric_column(df: pd.DataFrame, col: str) -> pd.DataFrame:
    """清洗数值列：转数值、去空、去零"""
    if col not in df.columns:
        return df
    df[col] = pd.to_numeric(df[col], errors="coerce")
    df = df[df[col].notna() & (df[col] != 0)]
    return df


def clean_date_column(df: pd.DataFrame, col: str) -> pd.DataFrame:
    """清洗日期列"""
    if col not in df.columns:
        return df
    df[col] = pd.to_datetime(df[col], errors="coerce")
    return df


def clean_amount_column(df: pd.DataFrame, col: str) -> pd.DataFrame:
    """清洗金额列：去逗号、去货币符号、转数值"""
    if col not in df.columns:
        return df
    df[col] = (
        df[col].astype(str)
        .str.replace(",", "", regex=False)
        .str.replace("¥", "", regex=False)
        .str.replace("￥", "", regex=False)
        .str.strip()
    )
    df["_金额数值"] = pd.to_numeric(df[col], errors="coerce").fillna(0.0)
    return df


# ==================== 分类汇总 ====================

def classify_positive_negative(
    df: pd.DataFrame,
    sku_col: str = "SKU编码",
    qty_col: str = "数量",
) -> Tuple[pd.DataFrame, pd.DataFrame]:
    """
    按正负数量分类汇总（蓝单/红单，销售出库/退货入库）。

    Returns:
        (positive_df, negative_df) 按 SKU 汇总后的 DataFrame
    """
    df["_qty_abs"] = df[qty_col].abs()

    positive = df[df[qty_col] > 0].groupby(sku_col)["_qty_abs"].sum().reset_index()
    positive.columns = [sku_col, "positive_qty"]

    negative = df[df[qty_col] < 0].groupby(sku_col)["_qty_abs"].sum().reset_index()
    negative.columns = [sku_col, "negative_qty"]

    return positive, negative


# ==================== SKU 匹配 ====================

def match_sku(
    sku: str,
    name_map: Dict[str, str],
    unit_map: Dict[str, str] = None,
    threshold: float = 0.92,
    unmatched_log: List[str] = None,
) -> Tuple[str, str, str]:
    """
    多级 SKU 匹配：精确 → 去前导零 → 补前导零 → 模糊匹配。

    Returns:
        (matched_sku, name, unit)
    """
    unit_map = unit_map or {}

    # 1. 精确匹配
    if sku in name_map:
        return sku, name_map[sku], unit_map.get(sku, "")

    # 2. 去前导零
    stripped = sku.lstrip("0")
    if stripped in name_map:
        return stripped, name_map[stripped], unit_map.get(stripped, "")

    # 3. 补前导零到 13 位
    padded = sku.zfill(13)
    if padded in name_map:
        return padded, name_map[padded], unit_map.get(padded, "")

    # 4. 模糊匹配
    best, best_ratio = None, 0
    for key in name_map:
        ratio = SequenceMatcher(None, sku, key).ratio()
        if ratio > best_ratio:
            best_ratio, best = ratio, key

    if best and best_ratio >= threshold:
        logger.debug(f"[模糊匹配] {sku} → {best} ({best_ratio:.0%})")
        return best, name_map[best], unit_map.get(best, "")

    if unmatched_log is not None:
        unmatched_log.append(sku)
    return sku, sku, ""


def load_material(file_path: str) -> Tuple[Dict[str, str], Dict[str, str]]:
    """
    加载物料表（第1列=SKU，第2列=名称，第3列=单位）。

    Returns:
        (name_map, unit_map)
    """
    if not os.path.exists(file_path):
        logger.warning(f"[data_tools] 物料表不存在: {file_path}")
        return {}, {}

    df = pd.read_excel(file_path, engine="openpyxl")
    name_map, unit_map = {}, {}

    for _, row in df.iterrows():
        sku = str(row.iloc[0]).strip()
        name = str(row.iloc[1]).strip() if pd.notna(row.iloc[1]) else ""
        unit = str(row.iloc[2]).strip() if len(row) > 2 and pd.notna(row.iloc[2]) else ""
        if sku and name:
            name_map[sku] = name
            unit_map[sku] = unit

    logger.info(f"[data_tools] 物料表加载: {len(name_map)} 条")
    return name_map, unit_map


# ==================== 日期工具 ====================

def parse_date(value: Any) -> Optional[datetime]:
    """把各种格式转成 datetime"""
    if value is None or isinstance(value, datetime):
        return value

    s = str(value).strip()
    if not s or s in ("nan", "None", "NaT", ""):
        return None

    s = s.split(" ")[0].split("T")[0]

    m = re.match(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", s)
    if m:
        try:
            return datetime(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            pass

    if s.isdigit() and len(s) == 8:
        try:
            return datetime(int(s[:4]), int(s[4:6]), int(s[6:8]))
        except ValueError:
            pass

    return None


def get_dominant_month_end(df: pd.DataFrame, date_col: str = "业务日期") -> str:
    """从数据中取出现最多的月份，返回该月最后一天"""
    dates = df[date_col].dropna() if date_col in df.columns else pd.Series(dtype=object)

    if len(dates) > 0:
        months = dates.apply(lambda x: (x.year, x.month))
        best = months.value_counts().idxmax()
        year, month = best
        last_day = calendar.monthrange(year, month)[1]
        return datetime(year, month, last_day).strftime("%Y-%m-%d")

    today = datetime.now()
    if today.day <= 10:
        return (today.replace(day=1) - timedelta(days=1)).strftime("%Y-%m-%d")
    return today.strftime("%Y-%m-%d")


# ==================== Excel 美化 ====================

def apply_border(filepath: str, row_count: int, col_count: int = 27) -> None:
    """给 Excel 加细边框（失败不影响主流程）"""
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Border, Side, Alignment

        wb = load_workbook(filepath)
        ws = wb.active
        thin = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin"),
        )

        for row in ws.iter_rows(max_row=row_count + 1, max_col=col_count):
            for cell in row:
                cell.border = thin
                cell.alignment = Alignment(
                    horizontal="left" if cell.row > 1 else "center",
                    vertical="center",
                )

        wb.save(filepath)
    except Exception:
        pass


# ==================== 结果输出 ====================

def save_output(
    df: pd.DataFrame,
    output_dir: str,
    prefix: str = "output",
    timestamp: str = None,
    **to_excel_kwargs,
) -> str:
    """统一保存输出文件"""
    os.makedirs(output_dir, exist_ok=True)
    ts = timestamp or datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = os.path.join(output_dir, f"{prefix}_{ts}.xlsx")
    df.to_excel(output_file, **to_excel_kwargs)
    logger.info(f"[data_tools] 输出: {output_file}")
    return output_file


def save_empty_result(output_dir: str, prefix: str = "无数据") -> str:
    """无数据时生成空结果"""
    os.makedirs(output_dir, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = os.path.join(output_dir, f"{prefix}_{ts}.xlsx")
    pd.DataFrame([{"状态": "无有效数据", "时间": datetime.now()}]).to_excel(output_file, index=False)
    return output_file